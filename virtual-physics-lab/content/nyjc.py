import math, os, random, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'tools'))
from vpl import *
from plans import place

rnd = random.Random(2026_08_24)
P = Paper('nyjc', 'Nanyang Junior College', 'NYJC', '24 Aug 2026', os.path.join(PRIV, 'src', 'nyjc.pdf'))
X = os.path.join(APP, 'dl')
g = 9.81

# ======================= Q1 T-shaped pendulum =======================
given = [(36, 1.685), (196, 2.373), (484, 3.418), (576, 3.782)]      # read from Fig. 1.4
a_true, b_true = 0.00388, 1.545
def T_of(x, M=0.05):
    return math.sqrt(b_true + a_true * (M / 0.05) * x * x)
xa, xc = 10.0, 18.0
ta = [round(20 * T_of(xa) * (1 + rnd.gauss(0, 0.003)), 2) for _ in range(2)]
tc = [round(20 * T_of(xc) * (1 + rnd.gauss(0, 0.003)), 2) for _ in range(2)]
Ta, Tc = hu(sum(ta) / 40, 3), hu(sum(tc) / 40, 3)
pts = given + [(xa ** 2, Ta ** 2), (xc ** 2, Tc ** 2)]
mA, cA, r2A = linfit([p[0] for p in pts], [p[1] for p in pts])
P.question('Q1', 'T-shaped pendulum and damped oscillations', 26, [2, 3, 4, 5, 6, 7, 8],
           'Masses on the cross-bar change the moment of inertia: T² = ax² + b. Add your own points to the printed graph, find a and b, '
           'test how they depend on M, plan how b depends on h and w, and analyse a velocity–time dataset of a damped oscillation.',
           sim='nyjc-q1', xlsx=[('NYJC_Q1_Tpendulum.xlsx', 'The four printed points plus two sample readings. Fit T² against x².'),
                                ('NYJC_Q1g_damped_dataset.xlsx', 'A practice version of datasetQ1g: horizontal velocity every 0.02 s. Find a (gradient) and s (cumulative area), plot a against s, and find ω, the amplitude and the damping ratio.')],
           skills=['metre rule', 'timing oscillations', 'plotting on a printed grid', 'numerical differentiation & integration', 'planning'])
P.ans(3, 'x =', f'{xa:.1f} cm', q='Q1')
P.text(3, 140, 452, 'Rule along cross-bar: nail centre 50.0 cm, bolt centre 60.0 cm', size=9, q='Q1')
P.text(3, 140, 560, f'Time for 20 oscillations: {ta[0]:.2f} s, {ta[1]:.2f} s\nmean = {sum(ta)/2:.3f} s, T = mean/20', size=9.5, q='Q1')
P.ans(3, 'T =', f'{Ta:.3f} s', q='Q1')
P.note(3, 460, 'x runs from the centre of the bolt to the nail (both centres). Measure on both sides and check they are equal.', 'tip')
P.note(3, 560, 'T ≈ 1.4 s: time 20 oscillations (≈ 28 s) twice; start counting at the centre of the swing using a fiducial mark.', 'mark')
grid = dict(rect=(167.0, 190.8, 497.5, 463.0), xr=[0, 600], yr=[1.5, 4.0])
P.add(4, dict(t='grid', **grid, pts=[(xa ** 2, Ta ** 2)], labels=[(xa ** 2 + 12, Ta ** 2 - 0.02, f'({xa**2:.0f}, {Ta**2:.2f})')]), q='Q1')
P.text(4, 300, 548, f'x² = {xa**2:.0f} cm², T² = ({Ta:.3f})² = {Ta**2:.3f} s²', size=9.5, q='Q1')
P.add(4, dict(t='grid', **grid, pts=[(xc ** 2, Tc ** 2)], labels=[(xc ** 2 + 12, Tc ** 2 - 0.02, f'({xc**2:.0f}, {Tc**2:.2f})')]), q='Q1')
P.text(4, 140, 640, f'x = {xc:.1f} cm (fills the gap between the printed points)\n20 oscillations: {tc[0]:.2f} s, {tc[1]:.2f} s → T = {Tc:.3f} s\nx² = {xc**2:.0f} cm², T² = {Tc**2:.3f} s²', size=9.2, q='Q1')
x1g, x2g = 40, 560
P.add(4, dict(t='grid', **grid, segs=[((0, cA), (600, cA + 600 * mA)), ((x1g, cA + mA * x1g), (x2g, cA + mA * x1g), '2,2'), ((x2g, cA + mA * x1g), (x2g, cA + mA * x2g), '2,2')],
              labels=[(330, cA + mA * x1g - 0.12, 'Δx² = 520 cm²'), (x2g - 140, cA + mA * 380, f'ΔT² = {mA*520:.2f} s²')]), q='Q1')
P.text(5, 120, 100, f'gradient = {mA*520:.3f} / 520 = {mA:.5f} s² cm⁻²\ny-intercept (read at x² = 0) = {cA:.2f} s²', size=9.5, q='Q1')
P.ans(5, 'a =', f'{mA:.5f}', q='Q1')
P.ans(5, 'b =', f'{cA:.2f}', q='Q1')
# (e) heavier masses
M2 = 0.100
te = []
for x in (10.0, 20.0):
    t20 = round(20 * T_of(x, M2) * (1 + rnd.gauss(0, 0.003)), 2)
    te.append((x, t20, hu(t20 / 20, 3)))
a2 = (te[1][2] ** 2 - te[0][2] ** 2) / (te[1][0] ** 2 - te[0][0] ** 2)
b2 = te[0][2] ** 2 - a2 * te[0][0] ** 2
P.table(5, 121, 318, ['M / g', 'x / cm', 't (20 osc) / s', 'T / s', 'T² / s²', 'x² / cm²'],
        [[100, f'{x:.1f}', f'{t:.2f}', f'{T:.3f}', f'{T*T:.3f}', f'{x*x:.0f}'] for x, t, T in te], size=8.8, q='Q1')
P.text(5, 121, 384, f'M = 100 g: a = ({te[1][2]**2:.3f} − {te[0][2]**2:.3f})/(400 − 100) = {a2:.5f} s² cm⁻²;  b = {b2:.2f} s²\n'
       f'M = 50 g (from (d)): a = {mA:.5f} s² cm⁻²;  b = {cA:.2f} s²', size=8.8, q='Q1')
P.write_lines(5, f'Doubling M doubles a (ratio {a2/mA:.2f} ≈ 2), so a is proportional to M. b is unchanged ({b2:.2f} s² vs {cA:.2f} s²), '
              'so b does not depend on M.', after=430, before=520, size=9.2, q='Q1')
P.note(4, 470, 'Plot your point with a neat cross at the exact position, and label it if the question asks.', 'mark')
P.note(4, 620, 'Choose x to fill the gap in the printed data (between x² ≈ 200 and 480 cm²) so your line is well defined.', 'tip')
P.note(5, 70, 'Use a large gradient triangle (> half the line) and read the intercept at x² = 0 (the x-axis starts at 0 here).', 'mark')
P.note(5, 300, '"Two further sets" with a different M: take two values of x so that a and b can be found for that M. Two points are enough for a straight line.', 'why')
P.note(5, 430, 'Conclusion: state how a depends on M (proportional) AND how b depends on M (independent), with numbers.', 'mark')
place(P, 'nyjc-q1f', [('box', 6, 119, 416, 410, 107, 13.5), ('lines', 6, 500, 740)], notes=[(6, 420), (6, 560)], q='Q1')

# (g) damped dataset
A0, w0, gam = 0.040, 2 * math.pi / 1.25, 0.050
dsg = []
for i in range(0, 501):
    t = round(i * 0.02, 2)
    x = A0 * math.exp(-gam * t) * math.cos(w0 * t)
    v = A0 * math.exp(-gam * t) * (-gam * math.cos(w0 * t) - w0 * math.sin(w0 * t))
    dsg.append((t, round(v + rnd.gauss(0, 0.0004), 4)))
acc, sdisp = [], [0.0]
for i in range(1, len(dsg)):
    sdisp.append(sdisp[-1] + 0.5 * (dsg[i][1] + dsg[i - 1][1]) * 0.02)
for i in range(1, len(dsg) - 1):
    acc.append(((dsg[i + 1][1] - dsg[i - 1][1]) / 0.04, sdisp[i]))
mg_, cg_, r2g = linfit([s for a, s in acc[:150]], [a for a, s in acc[:150]])
omega = math.sqrt(-mg_)
amp = abs(cg_ / mg_)
peaks = [abs(v) for v in [max(dsg[j][1] for j in range(k * 62, (k + 1) * 62)) for k in range(4)]]
ratio = math.exp(-gam * 1.25)
P.write_lines(7, '(1) a = gradient of the v–t graph: aₙ = (vₙ₊₁ − vₙ₋₁)/(tₙ₊₁ − tₙ₋₁) in a new column. '
              '(2) s = area under the v–t graph from t = 0: cumulative sum of ½(vₙ + vₙ₋₁)Δt (trapezium rule).', after=150, size=9.2, q='Q1')
P.graph(7, 140, 318, 300, 205, q='Q1', xlabel='s / m', ylabel='a / m s⁻²', xr=[-0.085, 0.005], yr=[-1.2, 1.2],
        series=[dict(pts=[(s, a) for a, s in acc[:150:4]], mark='dot')], lines=[dict(m=mg_, c=cg_)], eq=eqn(mg_, cg_, 4), eqpos=(290, 330), nx=4, ny=4)
P.ans(7, 'angular frequency', f'{omega:.2f} rad s⁻¹', q='Q1')
P.text(7, 140, 600, f'a = −ω²x = −ω²(s + A) ⇒ intercept = −ω²A\nA = intercept/gradient = {cg_:.3f}/{mg_:.2f}', size=9.3, q='Q1')
P.ans(7, 'initial amplitude', f'{amp*100:.1f} cm', q='Q1')
P.text(8, 140, 110, f'Peak speeds in successive cycles are proportional to the amplitudes:\n'
       f'v₁ = {peaks[0]:.4f}, v₂ = {peaks[1]:.4f}, v₃ = {peaks[2]:.4f}, v₄ = {peaks[3]:.4f} m s⁻¹\n'
       f'ratios: {peaks[1]/peaks[0]:.3f}, {peaks[2]/peaks[1]:.3f}, {peaks[3]/peaks[2]:.3f} → mean', size=9.2, q='Q1')
P.ans(8, '=', f'{(peaks[3]/peaks[0])**(1/3):.3f}', q='Q1')
P.note(7, 150, 'Name the method AND the spreadsheet operation: gradient (central difference) for a; cumulative area (trapezium) for s.', 'mark')
P.note(7, 330, 'For SHM a = −ω²x. Since s is measured from the release point, the graph of a against s has gradient −ω² but does not pass through the origin.', 'why')
P.note(7, 540, 'Use only the first few cycles for the fit: damping slowly shrinks the oscillation.', 'tip')
P.note(8, 100, 'Constant ratio A(n+1)/A(n) means exponential damping. Use peak speeds (or peak displacements) of consecutive cycles.', 'why')
workbook(os.path.join(X, 'NYJC_Q1_Tpendulum.xlsx'), 'NYJC Q1 – T-shaped pendulum (sample data)',
         'Four points read from Fig. 1.4 plus two sample measurements. Fit T² against x²: gradient = a, intercept = b.',
         ['x² / cm²', 'T² / s²'], [[round(a_, 1), round(b_, 3)] for a_, b_ in pts],
         chart=dict(x=1, y=2, title='T² against x²', xt='x² / cm²', yt='T² / s²'),
         results=[('a = gradient', '=SLOPE(B{first}:B{last},A{first}:A{last})', 's² cm⁻²'), ('b = intercept', '=INTERCEPT(B{first}:B{last},A{first}:A{last})', 's²')])
workbook(os.path.join(X, 'NYJC_Q1g_damped_dataset.xlsx'), 'NYJC Q1(g) – Damped oscillation velocity data (practice dataset)',
         'Δt = 0.02 s. Compute a (central difference) and s (cumulative trapezium), then plot a against s for the first 3 s.',
         ['t / s', 'v / m s⁻¹'], [list(r) for r in dsg],
         proc=[('s / m', '=IF(ROW()=6,0,G{r})'.replace('G{r}', 'C' + '{r}')),
               ('a / m s⁻²', '=IF(OR(ROW()=6,A{r1}=""),"",(B{r1}-B5)/0.04)')],
         answer_note='See the notes column: s is a running sum and a is a central difference.')
# replace processed formulas with proper ones (rows reference neighbours)
from openpyxl import load_workbook
fn = os.path.join(X, 'NYJC_Q1g_damped_dataset.xlsx')
wb = load_workbook(fn)
ws = wb['Suggested solution']
n = len(dsg)
for i in range(n):
    r = 6 + i
    ws.cell(r, 3).value = 0 if i == 0 else f'=C{r-1}+0.5*(B{r}+B{r-1})*0.02'
    ws.cell(r, 4).value = '' if i in (0, n - 1) else f'=(B{r+1}-B{r-1})/0.04'
ws['F5'], ws['G5'] = 'Result', 'Value'
ws['F6'], ws['G6'] = 'gradient of a–s (first 3 s)', f'=SLOPE(D7:D155,C7:C155)'
ws['F7'], ws['G7'] = 'ω = √(−gradient) / rad s⁻¹', '=SQRT(-G6)'
ws['F8'], ws['G8'] = 'A = intercept/gradient / m', '=INTERCEPT(D7:D155,C7:C155)/G6'
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.trendline import Trendline
ch = ScatterChart(); ch.title = 'a against s (first 3 s)'; ch.x_axis.title = 's / m'; ch.y_axis.title = 'a / m s⁻²'; ch.legend = None
ch.x_axis.delete = False; ch.y_axis.delete = False
sr = Series(Reference(ws, min_col=4, min_row=7, max_row=155), Reference(ws, min_col=3, min_row=7, max_row=155))
sr.marker.symbol = 'circle'; sr.marker.size = 3; sr.graphicalProperties.line.noFill = True
sr.trendline = Trendline(trendlineType='linear', dispEq=True, dispRSqr=True); ch.series.append(sr)
ws.add_chart(ch, 'F11'); wb.save(fn)
import base64
open(fn + '.b64.txt', 'w').write(base64.b64encode(open(fn, 'rb').read()).decode())

# ======================= Q2 friction on a leaning strip =======================
mu, S = 0.40, 0.100
def theta(M, d, L):
    return math.degrees(math.atan((M / 2 + S * d / L) / ((M + S) * mu)))
strips = {'long': dict(L=50.0, M=0.0452, d=10.0), 'short': dict(L=30.0, M=0.0271, d=10.0)}
res2 = {}
for k_, s in strips.items():
    L, M, d = s['L'], s['M'], s['d']
    thA = round(theta(M, d, L) + rnd.gauss(0, 0.7))
    thB = round(theta(M, L - d, L) + rnd.gauss(0, 0.7))
    FA = (M / 2 + S * d / L) / ((M + S) * math.tan(math.radians(thA)))
    FB = (M / 2 + S * (L - d) / L) / ((M + S) * math.tan(math.radians(thB)))
    res2[k_] = dict(L=L, M=M, dA=d, dB=L - d, thA=thA, thB=thB, FA=FA, FB=FB, k=FA / FB)
lo, sh = res2['long'], res2['short']
pct_th = 2 / lo['thA'] * 100
P.question('Q2', 'Friction on a strip leaning on a smooth board', 13, [9, 10, 11, 12, 13],
           'A strip leans on a smooth vertical board at the angle where it just stops slipping. F = (M/2 + Sd/L)/((M + S) tan θ) is the friction '
           'coefficient; compare F with the mass low (A) and high (B) for two strips.',
           sim='nyjc-q2', skills=['protractor', 'judging a limiting position', 'uncertainty', 'evaluation'])
P.ans(9, 'L =', f'{lo["L"]:.1f} cm', q='Q2')
P.ans(9, 'd A =', f'{lo["dA"]:.1f} cm', q='Q2')
P.note(9, 150, 'Strip lengths to 0.1 cm with a metre rule; d is measured to the centre of the slotted mass.', 'mark')
P.ans(10, 'A =', f'{lo["thA"]}°', q='Q2')
P.text(10, 140, 250, f'Repeats {lo["thA"]-1}°, {lo["thA"]}°, {lo["thA"]+1}°; the "just stays" position is hard to judge, so take\nΔθ ≈ 2° (larger than the half-range of 1°). % uncertainty = 2/{lo["thA"]} × 100%', size=9.2, q='Q2')
P.ans(10, 'percentage uncertainty', f'{pct_th:.0f}%', q='Q2')
P.ans(10, 'M =', f'{lo["M"]*1000:.1f} g', q='Q2')
P.text(10, 140, 560, f'F(A) = (0.0{lo["M"]*10000:.0f}/2 + 0.100 × {lo["dA"]:.1f}/{lo["L"]:.1f}) / ((0.0{lo["M"]*10000:.0f} + 0.100) tan {lo["thA"]}°)'.replace('0.0452', '0.0452'), size=9.2, q='Q2')
P.ans(10, 'F A =', f'{lo["FA"]:.2f}', q='Q2')
P.note(10, 200, 'Measure θ at the bottom of the strip with a protractor on the bench, or measure height and base and use tan θ = height/base.', 'tip')
P.note(10, 300, 'The limiting position varies between trials. Repeat it and use the spread for Δθ.', 'mark')
P.note(10, 600, 'F is dimensionless (it is the coefficient of static friction between strip and bench). Use consistent mass units.', 'why')
P.ans(11, 'd B =', f'{lo["dB"]:.1f} cm', q='Q2')
P.ans(11, 'B =', f'{lo["thB"]}°', after=500, q='Q2')
_xb, _yb = P.find(11, 'F B =')
P.text(11, 140, _yb - 26, f'F(B) = (0.0{lo["M"]*10000:.0f}/2 + 0.100 × {lo["dB"]:.1f}/{lo["L"]:.1f}) / ((0.0{lo["M"]*10000:.0f} + 0.100) tan {lo["thB"]}°)', size=9.2, q='Q2')
P.ans(11, 'F B =', f'{lo["FB"]:.2f}', q='Q2')
P.note(11, 430, 'With the mass high up, the strip needs a steeper angle before it can stand: θ_B > θ_A.', 'why')
for lab, val in [('L =', f'{sh["L"]:.1f} cm'), ('d A =', f'{sh["dA"]:.1f} cm'), ('M =', f'{sh["M"]*1000:.1f} g'),
                 ('F A =', f'{sh["FA"]:.2f}'), ('d B =', f'{sh["dB"]:.1f} cm'), ('F B =', f'{sh["FB"]:.2f}')]:
    P.ans(12, lab, val, q='Q2')
P.fill(12, 399, 189, f'{sh["thA"]}°', q='Q2')
P.fill(12, 399, 365, f'{sh["thB"]}°', q='Q2')
P.text(12, 120, 560, f'k = F(A) / F(B)\nlong strip: k = {lo["FA"]:.3f}/{lo["FB"]:.3f} = {lo["k"]:.2f}\nshort strip: k = {sh["FA"]:.3f}/{sh["FB"]:.3f} = {sh["k"]:.2f}', size=9.5, q='Q2')
P.ans(12, 'first value of k', f'{lo["k"]:.2f}', q='Q2')
P.ans(12, 'second value of k', f'{sh["k"]:.2f}', q='Q2')
kl, ks_ = round(lo['k'], 2), round(sh['k'], 2)
pdk = abs(kl - ks_) / min(kl, ks_) * 100
utan = lambda th: math.radians(2) / (math.sin(math.radians(th)) * math.cos(math.radians(th))) * 100
uk = utan(lo['thA']) + utan(lo['thB'])
P.write_lines(13, f'% difference in k = ({max(kl, ks_):.2f} − {min(kl, ks_):.2f})/{min(kl, ks_):.2f} = {pdk:.0f}%. k ∝ tan θB / tan θA and, with Δθ = 2°, '
              f'the % uncertainty in tan θ is Δθ/(sin θ cos θ) ≈ {utan(lo["thA"]):.0f}% at {lo["thA"]}° and {utan(lo["thB"]):.0f}% at {lo["thB"]}°, so k is uncertain by ≈ {uk:.0f}%. '
              f'{pdk:.0f}% {"<" if pdk <= uk else ">"} {uk:.0f}%, so the results {"support" if pdk <= uk else "do not support"} F(A)/F(B) = k.', after=170, before=260, size=8.6, q='Q2')
P.write_lines(13, 'The friction varies along the bench and the "just stays" position is hard to judge, so θ is hard to reproduce '
              '(the strip may stay at slightly different angles in repeat trials).', after=290, before=380, size=9.2, q='Q2')
P.write_lines(13, 'Tape a sheet of uniform sandpaper to the bench under the foot of the strip so the friction is the same in every trial, '
              'and repeat the limiting-angle measurement 5 times and take the average θ.', after=425, before=515, size=9.2, q='Q2')
P.note(13, 160, 'Use the % uncertainty you calculated in (a)(iii); compare it with the % difference in k.', 'mark')
P.note(13, 290, 'Avoid "human error". Name the specific difficulty: judging the limiting position, or friction varying along the bench.', 'warn')

# ======================= Q3 two-wire circuit =======================
rho, dF, dG, E3, r3 = 4.9e-7, 0.274, 0.376, 1.50, 0.30
AF, AG = math.pi * (dF / 2000) ** 2, math.pi * (dG / 2000) ** 2
def I_of(xcm):
    x = xcm / 100
    return E3 / (rho * x / AF + rho * (1 - x) / AG + r3)
rows3 = [(x, round(I_of(x) * (1 + rnd.gauss(0, 0.003)), 3)) for x in (10.0, 25.0, 40.0, 50.0, 65.0, 80.0, 90.0)]
m3, c3, r23 = linfit([x / 100 for x, I in rows3], [1 / I for x, I in rows3])
ratio3 = math.sqrt(1 + 1.000 * m3 / c3)
P.question('Q3', 'Two wires in series via a sliding lead', 11, [14, 15, 16],
           'Current flows along length x of wire F, through lead C, and along (L − x) of wire G. 1/I = Px + Q, and P/Q = (d_G²/d_F² − 1)/L '
           'lets you find the diameter ratio from the graph.',
           sim='nyjc-q3', xlsx=[('NYJC_Q3_two_wires.xlsx', 'Seven (x, I) pairs. Plot 1/I against x, find P and Q, then d_G/d_F.')],
           skills=['micrometer', 'ammeter readings', 'linearisation', 'resistivity reasoning'])
P.text(14, 140, 175, 'Zero error 0.00 mm; three readings along each wire, e.g. F: 0.27, 0.28, 0.27 mm;\nG: 0.38, 0.37, 0.38 mm', size=8.8, q='Q3')
P.ans(14, 'd F =', f'{dF:.2f} mm', q='Q3')
P.ans(14, 'd G =', f'{dG:.2f} mm', q='Q3')
P.ans(14, 'x =', '50.0 cm', q='Q3')
P.ans(14, 'I =', f'{dict(rows3)[50.0]:.3f} A', q='Q3')
P.note(14, 150, 'Average several micrometer readings at different points and orientations; check the zero error first.', 'mark')
P.note(14, 620, 'Open the switch between readings: the wire heats up and its resistance changes.', 'tip')
P.table(15, 160, 90, ['x / cm', 'I / A', '(1/I) / A⁻¹'], [[f'{x:.1f}', f'{I:.3f}', f'{1/I:.3f}'] for x, I in rows3], size=9.5, q='Q3')
P.graph(15, 140, 480, 300, 200, q='Q3', xlabel='x / m', ylabel='(1/I) / A⁻¹', xr=[0, 1.0], yr=[3.0, 6.0],
        series=[dict(pts=[(x / 100, 1 / I) for x, I in rows3])], lines=[dict(m=m3, c=c3)], eq=eqn(m3, c3, 4, r2=r23), eqpos=(330, 615))
P.ans(15, 'P =', f'{m3:.2f} A⁻¹ m⁻¹', q='Q3')
P.ans(15, 'Q =', f'{c3:.2f} A⁻¹', q='Q3')
P.text(16, 140, 200, f'P/Q = (d_G²/d_F² − 1)/L  ⇒  d_G/d_F = √(1 + LP/Q)\n= √(1 + 1.000 × {m3:.3f}/{c3:.3f})', size=9.8, q='Q3')
P.ans(16, '=', f'{ratio3:.2f}', q='Q3')
P.text(16, 140, 262, f'(The micrometer means 0.377/0.273 give 1.38. The graph value is a little\nlower because the cell\'s internal resistance adds to Q.)', size=8.8, q='Q3')
P.note(15, 90, 'Seven values over the whole wire (10–90 cm). Record I to the ammeter\'s precision, and 1/I to 3 or 4 s.f. consistently.', 'mark')
P.note(15, 480, 'Units: P in A⁻¹ m⁻¹ if x is in metres (A⁻¹ cm⁻¹ if in cm). Be consistent with L = 1.000 m in (d).', 'warn')
P.note(16, 190, 'Rearrange before substituting: d_G/d_F = √(1 + LP/Q).', 'mark')
workbook(os.path.join(X, 'NYJC_Q3_two_wires.xlsx'), 'NYJC Q3 – Two wires (sample data)', 'Plot 1/I against x (in m). P = gradient, Q = intercept.',
         ['x / cm', 'I / A'], [list(r) for r in rows3], proc=[('x / m', '=A{r}/100'), ('1/I / A⁻¹', '=1/B{r}')],
         chart=dict(x=3, y=4, title='1/I against x', xt='x / m', yt='1/I / A⁻¹'),
         results=[('P = gradient', '=SLOPE(D{first}:D{last},C{first}:C{last})', 'A⁻¹ m⁻¹'), ('Q = intercept', '=INTERCEPT(D{first}:D{last},C{first}:C{last})', 'A⁻¹'),
                  ('d_G/d_F = √(1 + LP/Q)', '=SQRT(1+SLOPE(D{first}:D{last},C{first}:C{last})/INTERCEPT(D{first}:D{last},C{first}:C{last}))', '')])
P.build()
